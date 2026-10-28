from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from datetime import date
import sqlite3
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = "change-this-secret-key"

DATABASE = "students.db"


def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def create_tables():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                roll_number TEXT NOT NULL UNIQUE,
                branch TEXT,
                email TEXT
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                date TEXT NOT NULL,
                status TEXT NOT NULL CHECK(status IN ('Present', 'Absent')),
                UNIQUE(student_id, date),
                FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE
            )
        """)


@app.route("/")
def index():
    today = date.today().isoformat()
    with get_db() as conn:
        total_students = conn.execute("SELECT COUNT(*) FROM students").fetchone()[0]
        present_today = conn.execute(
            "SELECT COUNT(*) FROM attendance WHERE date = ? AND status = 'Present'", (today,)
        ).fetchone()[0]
        absent_today = conn.execute(
            "SELECT COUNT(*) FROM attendance WHERE date = ? AND status = 'Absent'", (today,)
        ).fetchone()[0]
        marked_today = present_today + absent_today
        attendance_percentage = round((present_today / marked_today) * 100, 2) if marked_today else 0
        students = conn.execute("SELECT * FROM students ORDER BY name").fetchall()
    return render_template(
        "index.html",
        total_students=total_students,
        present_today=present_today,
        absent_today=absent_today,
        attendance_percentage=attendance_percentage,
        students=students
    )


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        roll_number = request.form.get("roll_number", "").strip()
        branch = request.form.get("branch", "").strip()
        email = request.form.get("email", "").strip()

        if not name or not roll_number:
            flash("Student name and roll number are required.", "error")
            return redirect(url_for("register"))

        try:
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO students (name, roll_number, branch, email) VALUES (?, ?, ?, ?)",
                    (name, roll_number, branch, email)
                )
            flash("Student registered successfully.", "success")
        except sqlite3.IntegrityError:
            flash("That roll number is already registered.", "error")
        return redirect(url_for("register"))

    with get_db() as conn:
        students = conn.execute("SELECT * FROM students ORDER BY name").fetchall()
    return render_template("register.html", students=students)


@app.route("/attendance", methods=["GET", "POST"])
def attendance():
    if request.method == "POST":
        student_id = request.form.get("student_id", "").strip()
        status = request.form.get("status", "").strip()
        attendance_date = request.form.get("attendance_date", "").strip() or date.today().isoformat()

        if not student_id or status not in ("Present", "Absent"):
            flash("Please select a student and a valid attendance status.", "error")
            return redirect(url_for("attendance"))

        try:
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO attendance (student_id, date, status) VALUES (?, ?, ?)",
                    (student_id, attendance_date, status)
                )
            flash("Attendance marked successfully.", "success")
        except sqlite3.IntegrityError:
            flash("Attendance has already been marked for this student on that date.", "error")
        return redirect(url_for("attendance"))

    with get_db() as conn:
        students = conn.execute("SELECT * FROM students ORDER BY name").fetchall()
    return render_template("attendance.html", students=students, today=date.today().isoformat())


@app.route("/records")
def records():
    with get_db() as conn:
        attendance_records = conn.execute("""
            SELECT attendance.id, students.name, students.roll_number,
                   attendance.date, attendance.status
            FROM attendance
            JOIN students ON students.id = attendance.student_id
            ORDER BY attendance.date DESC, students.name
        """).fetchall()
    return render_template("records.html", attendance_records=attendance_records)


@app.route("/student-report")
def student_report():
    with get_db() as conn:
        report = conn.execute("""
            SELECT s.id, s.name, s.roll_number,
                   COUNT(a.id) AS total_days,
                   SUM(CASE WHEN a.status = 'Present' THEN 1 ELSE 0 END) AS present_days,
                   SUM(CASE WHEN a.status = 'Absent' THEN 1 ELSE 0 END) AS absent_days
            FROM students s
            LEFT JOIN attendance a ON s.id = a.student_id
            GROUP BY s.id, s.name, s.roll_number
            ORDER BY s.name
        """).fetchall()

    rows = []
    for item in report:
        total = item["total_days"] or 0
        present = item["present_days"] or 0
        absent = item["absent_days"] or 0
        percentage = round((present / total) * 100, 2) if total else 0
        rows.append({
            "name": item["name"],
            "roll_number": item["roll_number"],
            "total_days": total,
            "present_days": present,
            "absent_days": absent,
            "percentage": percentage
        })
    return render_template("student_report.html", report=rows)


@app.route("/search-attendance", methods=["GET", "POST"])
def search_attendance():
    selected_date = request.values.get("attendance_date", "").strip()
    attendance_records = []
    if selected_date:
        with get_db() as conn:
            attendance_records = conn.execute("""
                SELECT students.name, students.roll_number, attendance.date, attendance.status
                FROM attendance
                JOIN students ON students.id = attendance.student_id
                WHERE attendance.date = ?
                ORDER BY students.name
            """, (selected_date,)).fetchall()
    return render_template(
        "search_attendance.html",
        attendance_records=attendance_records,
        selected_date=selected_date
    )


@app.route("/export-attendance")
def export_attendance():
    with get_db() as conn:
        rows = conn.execute("""
            SELECT students.name, students.roll_number, attendance.date, attendance.status
            FROM attendance
            JOIN students ON students.id = attendance.student_id
            ORDER BY attendance.date DESC, students.name
        """).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance Records"
    headers = ["Student Name", "Roll Number", "Date", "Status"]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2457A7")
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append([row["name"], row["roll_number"], row["date"], row["status"]])

    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max_length + 3, 35)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="attendance_records.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    create_tables()
    app.run(debug=True)
